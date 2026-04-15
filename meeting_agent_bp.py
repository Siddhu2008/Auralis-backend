"""
Auralis AI Meeting Agent — REST API Blueprint
Provides endpoints for deploying, querying, and managing the AI meeting agent.
"""
from flask import Blueprint, request, jsonify
from utils.jwt_handler import decode_token
from utils.meeting_agent import (
    generate_meeting_report,
    answer_post_meeting_question,
    draft_email_from_meeting,
    detect_qa_pairs,
    generate_video_summary_script,
)
from utils.tts_handler import text_to_speech_base64
from models.meeting_qa import save_qa_pair, get_room_qa
import io
from docx import Document
from datetime import datetime as dt

meeting_agent_bp = Blueprint('meeting_agent', __name__)

# In-memory registry of active agents per room
_active_agents = {}  # room_id -> {user_id, deployed_at, transcript, qa_pairs}


def _auth(req):
    """Extract user_id from Authorization header."""
    auth_header = req.headers.get('Authorization')
    if not auth_header or not auth_header.startswith('Bearer '):
        return None
    token = auth_header.split(' ')[1]
    try:
        payload = decode_token(token)
        return payload['user_id']
    except Exception:
        return None


@meeting_agent_bp.route('/deploy', methods=['POST'])
def deploy_agent():
    """Deploy the AI agent to a meeting room."""
    user_id = _auth(request)
    if not user_id:
        return jsonify({"error": "Unauthorized"}), 401

    data = request.get_json(silent=True) or {}
    room_id = (data.get('room_id') or '').strip()
    if not room_id:
        return jsonify({"error": "room_id is required"}), 400

    from datetime import datetime
    _active_agents[room_id] = {
        'user_id': user_id,
        'deployed_at': datetime.utcnow().isoformat(),
        'transcript': [],
        'qa_pairs': [],
    }

    return jsonify({
        "status": "deployed",
        "room_id": room_id,
        "message": "Auralis AI Agent is now active in the meeting.",
    }), 200


@meeting_agent_bp.route('/retire', methods=['POST'])
def retire_agent():
    """Remove the AI agent from a meeting room."""
    user_id = _auth(request)
    if not user_id:
        return jsonify({"error": "Unauthorized"}), 401

    data = request.get_json(silent=True) or {}
    room_id = (data.get('room_id') or '').strip()

    if room_id in _active_agents:
        del _active_agents[room_id]

    return jsonify({"status": "retired", "room_id": room_id}), 200


@meeting_agent_bp.route('/status/<room_id>', methods=['GET'])
def agent_status(room_id):
    """Check if agent is active in a room."""
    user_id = _auth(request)
    if not user_id:
        return jsonify({"error": "Unauthorized"}), 401

    agent = _active_agents.get(room_id)
    if agent:
        return jsonify({
            "active": True,
            "deployed_at": agent['deployed_at'],
            "qa_count": len(agent.get('qa_pairs', [])),
            "transcript_lines": len(agent.get('transcript', [])),
        }), 200

    return jsonify({"active": False}), 200


@meeting_agent_bp.route('/ask', methods=['POST'])
def ask_agent():
    """Ask a question about a past meeting (post-meeting RAG)."""
    user_id = _auth(request)
    if not user_id:
        return jsonify({"error": "Unauthorized"}), 401

    data = request.get_json(silent=True) or {}
    question = (data.get('question') or '').strip()
    room_id = (data.get('room_id') or '').strip()

    if not question:
        return jsonify({"error": "question is required"}), 400

    # Get transcript from active agent or database
    transcript = ''
    qa_pairs = []
    agent = _active_agents.get(room_id)
    if agent:
        # Convert structured transcript to plain text for RAG
        t_data = agent.get('transcript', [])
        transcript = '\n'.join([f"{l['speaker']}: {l['text']}" for l in t_data])
        qa_pairs = agent.get('qa_pairs', [])
    else:
        # First check the Meeting model for persistent data
        try:
            from models.meeting import Meeting
            meeting = Meeting.query.filter_by(room_id=room_id, user_id=user_id).first()
            if meeting:
                if meeting.qa_pairs:
                    qa_pairs = meeting.qa_pairs
                if meeting.transcript:
                    transcript = meeting.transcript
        except Exception:
            pass

        # Fallback: use old meeting_qa table
        if not qa_pairs:
            qa_pairs = get_room_qa(room_id)

    # If no transcript from memory, try to get from meetings DB (legacy)
    if not transcript:
        try:
            from models.meeting import get_user_meetings
            meetings = get_user_meetings(user_id)
            for m in meetings:
                if m.get('room_id') == room_id and m.get('transcript'):
                    transcript = m['transcript']
                    break
        except Exception:
            pass

    answer = answer_post_meeting_question(question, transcript, qa_pairs)

    # Generate voice for the answer
    audio_b64 = text_to_speech_base64(answer)

    return jsonify({
        "answer": answer,
        "audio": audio_b64,
        "source": "meeting_memory",
    }), 200


@meeting_agent_bp.route('/draft-email', methods=['POST'])
def draft_email():
    """Draft or send an email based on meeting context."""
    user_id = _auth(request)
    if not user_id:
        return jsonify({"error": "Unauthorized"}), 401

    data = request.get_json(silent=True) or {}
    instruction = (data.get('instruction') or '').strip()
    room_id = (data.get('room_id') or '').strip()
    send_immediately = data.get('send', False)

    if not instruction:
        return jsonify({"error": "instruction is required"}), 400

    # Get meeting context
    transcript = ''
    qa_pairs = []
    agent = _active_agents.get(room_id)
    if agent:
        # Convert structured transcript to plain text for drafting
        t_data = agent.get('transcript', [])
        transcript = '\n'.join([f"{l['speaker']}: {l['text']}" for l in t_data])
        qa_pairs = agent.get('qa_pairs', [])

    draft = draft_email_from_meeting(instruction, transcript, qa_pairs)

    # Optionally send immediately
    if send_immediately and draft.get('to') and draft['to'] != 'unspecified@pending.com':
        try:
            from utils.email_handler import send_email_custom
            send_email_custom(draft['to'], draft['subject'], draft['body'])
            draft['sent'] = True
        except Exception as e:
            draft['sent'] = False
            draft['send_error'] = str(e)
    else:
        draft['sent'] = False

    return jsonify({"email": draft}), 200


@meeting_agent_bp.route('/report/<room_id>', methods=['GET'])
def get_report(room_id):
    """Get the structured meeting report for a room."""
    user_id = _auth(request)
    if not user_id:
        return jsonify({"error": "Unauthorized"}), 401

    # Get transcript
    transcript_data = []
    qa_pairs = []
    agent = _active_agents.get(room_id)
    title = f'Meeting {room_id}'

    report = ""
    if agent:
        transcript_data = agent.get('transcript', [])
        qa_pairs = agent.get('qa_pairs', [])
        participants = list(set([l['speaker'] for l in transcript_data if 'speaker' in l]))
        report = generate_meeting_report(transcript_data, qa_pairs, title, participants)
    else:
        # Check database for persistent report
        from models.meeting import Meeting
        meeting = Meeting.query.filter_by(room_id=room_id, user_id=user_id).first()
        if meeting and meeting.agent_report:
            report = meeting.agent_report
        else:
            # Legacy fallback / reconstruct from transcript
            qa_pairs = get_room_qa(room_id)
            if meeting:
                title = meeting.title or title
                t_text = meeting.transcript or ''
                if t_text.startswith('['):
                    try:
                        import json
                        transcript_data = json.loads(t_text)
                    except:
                        transcript_data = [{'speaker': 'System', 'text': t_text, 'timestamp': 'N/A'}]
                else:
                    transcript_data = [{'speaker': 'System', 'text': t_text, 'timestamp': 'N/A'}]
                
                if transcript_data:
                    participants = list(set([l['speaker'] for l in transcript_data if 'speaker' in l]))
                    report = generate_meeting_report(transcript_data, qa_pairs, title, participants)

    if not report:
        return jsonify({"error": "No report found for this room"}), 404

    return jsonify({"report": report, "room_id": room_id}), 200


@meeting_agent_bp.route('/export/doc/<room_id>', methods=['GET'])
def export_doc(room_id):
    """Export the meeting report as a DOCX file."""
    user_id = _auth(request)
    if not user_id:
        return jsonify({"error": "Unauthorized"}), 401

    # Get data (reusing logic from report)
    report_md = ""
    agent = _active_agents.get(room_id)
    transcript_data = []
    qa_pairs = []
    title = "Meeting Report"

    if agent:
        transcript_data = agent.get('transcript', [])
        qa_pairs = agent.get('qa_pairs', [])
        participants = list(set([l['speaker'] for l in transcript_data if 'speaker' in l]))
        report_md = generate_meeting_report(transcript_data, qa_pairs, title, participants)
    else:
        from models.meeting import Meeting
        m = Meeting.query.filter_by(room_id=room_id, user_id=user_id).first()
        if m and m.agent_report:
            report_md = m.agent_report
            title = m.title or title # Ensure title is set from meeting if report exists
        else:
            # Fallback reconstruction
            qa_pairs = get_room_qa(room_id)
            if m:
                title = m.title or title
                try:
                    import json
                    transcript_data = json.loads(m.transcript)
                except:
                    transcript_data = [{'speaker': 'System', 'text': m.transcript, 'timestamp': 'N/A'}]
                
                if transcript_data:
                    participants = list(set([l['speaker'] for l in transcript_data if 'speaker' in l]))
                    report_md = generate_meeting_report(transcript_data, qa_pairs, title, participants)

    if not report_md:
        return jsonify({"error": "No data found"}), 404

    # Convert MD-like sections to DOCX
    doc = Document()
    doc.add_heading(f'Meeting Report: {title}', 0)

    # Simplified parser for our specific report structure
    sections = report_md.split('## ')
    for sect in sections:
        if not sect.strip(): continue
        lines = sect.split('\n')
        header = lines[0].strip()
        doc.add_heading(header, level=1)
        content = '\n'.join(lines[1:]).strip()
        doc.add_paragraph(content)

    doc_io = io.BytesIO()
    doc.save(doc_io)
    doc_io.seek(0)

    from flask import send_file
    return send_file(
        doc_io,
        as_attachment=True,
        download_name=f"Auralis_Report_{room_id}.docx",
        mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    )


@meeting_agent_bp.route('/report/video/<room_id>', methods=['GET'])
def get_video_script(room_id):
    """Get a video summary script for a meeting."""
    user_id = _auth(request)
    if not user_id:
        return jsonify({"error": "Unauthorized"}), 401

    agent = _active_agents.get(room_id)
    transcript_data = []
    title = f'Meeting {room_id}'

    if agent:
        transcript_data = agent.get('transcript', [])
    else:
        from models.meeting import Meeting
        m = Meeting.query.filter_by(room_id=room_id, user_id=user_id).first()
        if m:
            title = m.title
            try:
                import json
                transcript_data = json.loads(m.transcript)
            except:
                transcript_data = [{'speaker': 'System', 'text': m.transcript, 'timestamp': 'N/A'}]

    if not transcript_data:
        return jsonify({"error": "No data found"}), 404

    script = generate_video_summary_script(transcript_data, title)
    return jsonify({"script": script, "room_id": room_id}), 200


@meeting_agent_bp.route('/qa/<room_id>', methods=['GET'])
def get_qa(room_id):
    """Get all Q&A pairs from a meeting."""
    user_id = _auth(request)
    if not user_id:
        return jsonify({"error": "Unauthorized"}), 401

    # Check in-memory first
    agent = _active_agents.get(room_id)
    if agent:
        return jsonify({"qa_pairs": agent.get('qa_pairs', []), "source": "live"}), 200

    # Fallback to database
    qa_pairs = get_room_qa(room_id)
    return jsonify({"qa_pairs": qa_pairs, "source": "database"}), 200


# ─────────────────────────────────────────────────────────────────────────────
# Helper for socket_events to interact with the agent
# ─────────────────────────────────────────────────────────────────────────────
def is_agent_active(room_id):
    return room_id in _active_agents


def feed_transcript_to_agent(room_id, text, user_id=None):
    """Feed a transcript line to the agent for Q&A detection."""
    agent = _active_agents.get(room_id)
    if not agent:
        return None

    # Get metadata for the transcript entry
    # Note: In socket_events.py we already store entries in live_transcripts.
    # Here we should also keep agent['transcript'] in sync if it's used for reporting.
    # We'll detect the current timestamp and speaker if possible, or just use what's passed.
    timestamp = dt.utcnow().strftime('%H:%M:%S')
    
    # We'll assume 'text' is the raw string for now.
    # Better yet, let's keep agent['transcript'] as the source of truth for the agent.
    entry = {'speaker': 'Participant', 'text': text, 'timestamp': timestamp}
    agent['transcript'].append(entry)

    # Detect Q&A using ML model
    from utils.meeting_agent import detect_qa
    label, confidence = detect_qa(text)

    if label == 'question' and confidence > 0.6:
        agent['_pending_question'] = text
    elif label == 'answer' and confidence > 0.6 and agent.get('_pending_question'):
        qa_pair = {
            'question': agent['_pending_question'],
            'answer': text,
            'speaker': 'Participant',
        }
        agent['qa_pairs'].append(qa_pair)
        agent.pop('_pending_question', None)

        # Save to database
        if user_id:
            try:
                save_qa_pair(room_id, user_id, qa_pair['question'], qa_pair['answer'])
            except Exception:
                pass

        return qa_pair

    return None


def finalize_agent_meeting(room_id, user_id=None, title='Meeting'):
    """Called when meeting ends — generates report and persists Q&A."""
    agent = _active_agents.get(room_id)
    if not agent:
        return None

    transcript_data = agent.get('transcript', [])
    qa_pairs = agent.get('qa_pairs', [])

    # Extract participants
    participants = list(set([l['speaker'] for l in transcript_data if 'speaker' in l]))

    # Save any remaining Q&A pairs to database
    if user_id:
        for qa in qa_pairs:
            try:
                save_qa_pair(room_id, user_id, qa['question'], qa['answer'])
            except Exception:
                pass

    # Generate report
    report = generate_meeting_report(transcript_data, qa_pairs, title, participants)

    # Index into vector store for continuous learning
    try:
        from utils.vector_store import vector_store
        import json
        t_text = json.dumps(transcript_data)
        combined_text = f"Title: {title}\n\nTranscript:\n{t_text}"
        if qa_pairs:
            qa_text = '\n'.join(f"Q: {p['question']} A: {p['answer']}" for p in qa_pairs)
            combined_text += f"\n\nQ&A:\n{qa_text}"
        vector_store.add_meeting(room_id, combined_text, metadata={'title': title})
    except Exception:
        pass

    # Cleanup
    _active_agents.pop(room_id, None)

    return report, qa_pairs
